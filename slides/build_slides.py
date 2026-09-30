"""python -m slides [--results results] [--out slides/compass.pptx] : the 30-minute deck.

Every number and chart comes from the analysis outputs; the outline holds words only."""

from __future__ import annotations

import argparse
import os
import sys

import pandas as pd
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

TEXT = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)


def load_outline(path: str) -> dict:
    with open(path) as f:
        return yaml.safe_load(f)


def _title(slide, text: str):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = text
    p.font.size, p.font.bold, p.font.color.rgb = Pt(26), True, TEXT


def add_title_slide(prs, title: str, subtitle: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(12), Inches(1.5))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(34), True, TEXT
    p2 = tb.text_frame.add_paragraph()
    p2.text = subtitle
    p2.font.size, p2.font.color.rgb = Pt(18), MUTED
    return s


def add_bullets(
    slide, bullets: list[str], left: float, top: float, width: float, height: float, size: int = 14
):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + b
        p.font.size, p.font.color.rgb = Pt(size), TEXT
        p.space_after = Pt(6)


def add_table(
    slide, df: pd.DataFrame, left: float, top: float, width: float, max_rows: int = 14, font: int = 10
):
    df = df.head(max_rows)
    cols = [c for c in df.columns if c not in ("bytes",)]
    rows, ncols = len(df) + 1, len(cols)
    shape = slide.shapes.add_table(rows, ncols, Inches(left), Inches(top), Inches(width), Inches(0.3 * rows))
    t = shape.table
    for j, c in enumerate(cols):
        cell = t.cell(0, j)
        cell.text = str(c)
        cell.text_frame.paragraphs[0].font.size, cell.text_frame.paragraphs[0].font.bold = Pt(font), True
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, c in enumerate(cols):
            v = r[c]
            cell = t.cell(i, j)
            cell.text = f"{v:.3g}" if isinstance(v, float) else str(v)
            cell.text_frame.paragraphs[0].font.size = Pt(font)


def add_section(prs, sec: dict, results_dir: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, sec["title"])
    chart = sec.get("chart")
    chart_path = os.path.join(results_dir, "charts", f"{chart}.png") if chart else None
    have_chart = bool(chart_path and os.path.exists(chart_path))
    table = sec.get("table")
    table_path = os.path.join(results_dir, "tables", f"{table}.csv") if table else None
    df = pd.read_csv(table_path) if table_path and os.path.exists(table_path) else None
    bullets = sec.get("bullets", [])
    if have_chart:
        s.shapes.add_picture(chart_path, Inches(0.5), Inches(1.3), width=Inches(7.6))
        if bullets:
            add_bullets(s, bullets, 8.3, 1.3, 4.6, 5.5, size=12)
        if df is not None:
            add_table(s, df, 0.5, 5.6, 7.6, max_rows=5, font=8)
    else:
        if bullets:
            add_bullets(s, bullets, 0.6, 1.3, 12.0 if df is None else 6.0, 5.5)
        if df is not None:
            add_table(s, df, 6.8 if bullets else 0.6, 1.3, 6.0 if bullets else 12.0, max_rows=16, font=9)
    if not have_chart and chart:
        tb = s.shapes.add_textbox(Inches(0.6), Inches(6.9), Inches(12), Inches(0.4))
        tb.text_frame.paragraphs[0].text = f"(chart {chart} not yet produced — run `make analyze`)"
        tb.text_frame.paragraphs[0].font.size, tb.text_frame.paragraphs[0].font.color.rgb = Pt(10), MUTED
    return s


def build(outline_path: str, results_dir: str, out: str) -> str:
    o = load_outline(outline_path)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    add_title_slide(prs, o["title"], o.get("subtitle", ""))
    for sec in o["sections"]:
        add_section(prs, sec, results_dir)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    prs.save(out)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--outline", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "outline.yaml")
    )
    ap.add_argument("--results", default="results")
    ap.add_argument("--out", default="slides/compass.pptx")
    a = ap.parse_args(argv)
    print(build(a.outline, a.results, a.out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
