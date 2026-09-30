from pptx import Presentation

from analysis.__main__ import main as analysis_main
from slides.build_slides import build, load_outline


def test_deck_has_one_slide_per_section_plus_title(fixtures, tmp_path):
    assert (
        analysis_main(
            [str(fixtures / "results"), "--out", str(tmp_path / "res"), "--readme", str(tmp_path / "none.md")]
        )
        == 0
    )
    outline = "slides/outline.yaml"
    out = build(outline, str(tmp_path / "res"), str(tmp_path / "deck.pptx"))
    prs = Presentation(out)
    assert len(prs.slides) == 1 + len(load_outline(outline)["sections"])
    # the chart slides embed a picture
    pics = sum(1 for s in prs.slides for sh in s.shapes if sh.shape_type == 13)
    assert pics >= 4
